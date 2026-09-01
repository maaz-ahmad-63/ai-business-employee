import hashlib
import io
import mimetypes
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, BinaryIO, Dict, List, Optional, Union


@dataclass
class DocumentSection:
    """A section of extracted document text with location/source metadata."""
    text: str
    metadata: Dict[str, Any] = field(default_factory=dict)


@dataclass
class ParsedDocument:
    """Result of parsing a document."""
    filename: str
    mime_type: str
    file_size: int
    checksum: str
    sections: List[DocumentSection]
    metadata: Dict[str, Any] = field(default_factory=dict)

    @property
    def full_text(self) -> str:
        return "\n\n".join(s.text for s in self.sections if s.text.strip())


class DocumentParser:
    """Parses text, PDF, and DOCX documents with metadata preservation."""

    @staticmethod
    def calculate_checksum(data: bytes) -> str:
        return hashlib.sha256(data).hexdigest()

    @classmethod
    def parse_bytes(
        cls,
        data: bytes,
        filename: str,
        mime_type: Optional[str] = None,
        custom_metadata: Optional[Dict[str, Any]] = None,
    ) -> ParsedDocument:
        """Parse in-memory file bytes."""
        file_size = len(data)
        checksum = cls.calculate_checksum(data)

        if not mime_type:
            mime_type, _ = mimetypes.guess_type(filename)
            mime_type = mime_type or "application/octet-stream"

        ext = Path(filename).suffix.lower()
        sections: List[DocumentSection] = []

        if ext in (".pdf",) or mime_type == "application/pdf":
            sections = cls._parse_pdf(data, filename)
        elif ext in (".docx",) or mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            sections = cls._parse_docx(data, filename)
        else:
            # Default to text / markdown
            sections = cls._parse_text(data, filename)

        doc_meta = {
            "filename": filename,
            "mime_type": mime_type,
            "file_size": file_size,
            "checksum": checksum,
            "section_count": len(sections),
            **(custom_metadata or {}),
        }

        return ParsedDocument(
            filename=filename,
            mime_type=mime_type,
            file_size=file_size,
            checksum=checksum,
            sections=sections,
            metadata=doc_meta,
        )

    @classmethod
    def parse_file(
        cls,
        file_path: Union[str, Path],
        custom_metadata: Optional[Dict[str, Any]] = None,
    ) -> ParsedDocument:
        """Parse file from disk path."""
        path = Path(file_path)
        if not path.is_file():
            raise FileNotFoundError(f"File not found: {file_path}")

        with open(path, "rb") as f:
            data = f.read()

        return cls.parse_bytes(
            data=data,
            filename=path.name,
            custom_metadata=custom_metadata,
        )

    @classmethod
    def _parse_text(cls, data: bytes, filename: str) -> List[DocumentSection]:
        """Extract text from UTF-8/plain text."""
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = data.decode("latin-1")
            except Exception:
                text = data.decode("utf-8", errors="replace")

        return [DocumentSection(text=text, metadata={"filename": filename})]

    @classmethod
    def _parse_pdf(cls, data: bytes, filename: str) -> List[DocumentSection]:
        """Extract text from PDF pages using pypdf."""
        import pypdf

        sections: List[DocumentSection] = []
        reader = pypdf.PdfReader(io.BytesIO(data))

        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                sections.append(
                    DocumentSection(
                        text=page_text,
                        metadata={
                            "filename": filename,
                            "page": idx + 1,
                            "total_pages": len(reader.pages),
                        },
                    )
                )

        return sections

    @classmethod
    def _parse_docx(cls, data: bytes, filename: str) -> List[DocumentSection]:
        """Extract text from DOCX paragraphs and tables using python-docx."""
        import docx

        doc = docx.Document(io.BytesIO(data))
        paragraphs: List[str] = []

        for p in doc.paragraphs:
            if p.text.strip():
                paragraphs.append(p.text)

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n\n".join(paragraphs)
        return [DocumentSection(text=full_text, metadata={"filename": filename})]
