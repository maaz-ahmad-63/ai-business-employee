import io
import pytest
from sqlalchemy import text

from app.ingestion.parser import DocumentParser
from app.ingestion.service import IngestionService
from app.storage.database import SessionLocal


def test_parse_text_bytes():
    raw_bytes = "This is a simple test document for parser validation.".encode("utf-8")
    parsed = DocumentParser.parse_bytes(raw_bytes, "test.txt", "text/plain")

    assert parsed.filename == "test.txt"
    assert parsed.file_size == len(raw_bytes)
    assert len(parsed.sections) == 1
    assert "simple test document" in parsed.sections[0].text


def test_parse_pdf_bytes():
    import pypdf

    # Create in-memory PDF
    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    # We can write a minimal PDF stream or use pypdf
    stream = io.BytesIO()
    writer.write(stream)
    pdf_bytes = stream.getvalue()

    parsed = DocumentParser.parse_bytes(pdf_bytes, "sample.pdf", "application/pdf")
    assert parsed.filename == "sample.pdf"
    assert parsed.mime_type == "application/pdf"


def test_parse_docx_bytes():
    import docx

    doc = docx.Document()
    doc.add_paragraph("First paragraph in test document.")
    doc.add_paragraph("Second paragraph with technical details.")
    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    parsed = DocumentParser.parse_bytes(docx_bytes, "sample.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert parsed.filename == "sample.docx"
    assert len(parsed.sections) == 1
    assert "First paragraph" in parsed.sections[0].text
    assert "Second paragraph" in parsed.sections[0].text


def test_ingestion_service_text_end_to_end(test_tenant, db_session):
    svc = IngestionService()
    content = "The enterprise RAG platform provides high-throughput vector search for multi-tenant organizations."

    result = svc.ingest_text(
        tenant_id=test_tenant,
        text_content=content,
        filename="rag_spec.txt",
    )

    assert result["status"] == "completed"
    assert result["chunk_count"] >= 1
    doc_id = result["document_id"]

    # Verify in database
    doc_row = db_session.execute(
        text("SELECT status, filename FROM documents WHERE id = :id AND tenant_id = :tenant_id"),
        {"id": doc_id, "tenant_id": test_tenant},
    ).mappings().first()
    assert doc_row["status"] == "completed"
    assert doc_row["filename"] == "rag_spec.txt"

    chunk_rows = db_session.execute(
        text("SELECT count(*) as count FROM chunks WHERE document_id = :doc_id AND tenant_id = :tenant_id"),
        {"doc_id": doc_id, "tenant_id": test_tenant},
    ).mappings().first()
    assert chunk_rows["count"] >= 1
